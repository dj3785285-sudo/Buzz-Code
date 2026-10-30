"""OTT Title Matcher engine (single-file bundle of ott_core)."""
from __future__ import annotations

# ======================================================================
# module: errors
# ======================================================================
"""Exception types used across the application."""



class OTTError(Exception):
    """Base class for all application errors that should be shown to the user."""


class WorkbookError(OTTError):
    """The uploaded file is not a readable .xlsx workbook."""


class ColumnSpecError(OTTError, ValueError):
    """A column specification such as 'BK:BQ' could not be parsed."""


class MappingError(OTTError):
    """The question/column mapping is missing or invalid.

    ``errors`` holds one human-readable message per problem.
    """

    def __init__(self, errors):
        if isinstance(errors, str):
            errors = [errors]
        self.errors = list(errors)
        super().__init__("; ".join(self.errors))

# ======================================================================
# module: excel_utils
# ======================================================================
"""Small helpers for Excel column notation."""


import re
from typing import Iterable, List

from openpyxl.utils import column_index_from_string, get_column_letter



MAX_EXCEL_COLUMN = 16384  # XFD


def col_letter(index: int) -> str:
    return get_column_letter(index)


def addr(row: int, col: int) -> str:
    """1-based (row, col) -> 'BY19'."""
    return f"{get_column_letter(col)}{row}"


def col_to_index(token: str) -> int:
    t = token.strip().replace("$", "")
    if not re.fullmatch(r"[A-Za-z]{1,3}", t):
        raise ColumnSpecError(
            f"'{token}' is not a valid column. Use column letters only, e.g. BK or BK:BQ."
        )
    try:
        idx = column_index_from_string(t.upper())
    except ValueError as exc:  # pragma: no cover - defensive
        raise ColumnSpecError(f"'{token}' is not a valid Excel column.") from exc
    if idx > MAX_EXCEL_COLUMN:
        raise ColumnSpecError(f"Column '{token}' is beyond Excel's last column (XFD).")
    return idx


def parse_columns(spec: str | None) -> List[int]:
    """Parse 'BK:BQ', 'BK,BL,BM' or a mix ('BK:BM, BO') into 1-based column indexes.

    Order of first appearance is kept and duplicates are removed.
    """
    if spec is None:
        return []
    tokens = [t for t in re.split(r"[,\s;]+", str(spec).strip()) if t]
    out: List[int] = []
    seen = set()
    for tok in tokens:
        if ":" in tok:
            a, _, b = tok.partition(":")
            ia, ib = col_to_index(a), col_to_index(b)
            if ia > ib:
                ia, ib = ib, ia
            cols: Iterable[int] = range(ia, ib + 1)
        else:
            cols = [col_to_index(tok)]
        for c in cols:
            if c not in seen:
                seen.add(c)
                out.append(c)
    return out


def columns_to_spec(columns: Iterable[int]) -> str:
    """[63, 64, 65, 67] -> 'BK:BM,BO'."""
    cols = sorted(set(columns))
    if not cols:
        return ""
    parts = []
    start = prev = cols[0]
    for c in cols[1:]:
        if c == prev + 1:
            prev = c
            continue
        parts.append((start, prev))
        start = prev = c
    parts.append((start, prev))
    return ",".join(
        get_column_letter(a) if a == b else f"{get_column_letter(a)}:{get_column_letter(b)}"
        for a, b in parts
    )

# ======================================================================
# module: normalize
# ======================================================================
"""Text normalisation for OTT titles.

``normalize_title`` produces a canonical, space-separated lowercase form.
``compact_key`` additionally removes spaces so that joined/separated words
("BiggBoss" / "Bigg Boss") become identical.
"""


import re
import unicodedata
from typing import Tuple

_APOSTROPHES = dict.fromkeys(map(ord, "'’‘`´"), None)
_ORDINAL_SEASON = re.compile(r"\b(\d+)(?:st|nd|rd|th)\s+season\b")
_SEASON_WORD = re.compile(r"\b(?:seasons?|ssn|sn|s)\s*0*(\d+(?:\.\d+)?)\b")
_TRAILING_NUMBER = re.compile(r"\s*\d+(?:\.\d+)?$")
_DIGITS = re.compile(r"\d+(?:\.\d+)?")


def _strip_latin_accents(text: str) -> str:
    """Remove accents from Latin letters only (keeps Devanagari etc. intact)."""
    out = []
    prev = ""
    for ch in unicodedata.normalize("NFKD", text):
        if unicodedata.combining(ch) and prev and prev.isascii() and prev.isalpha():
            continue
        out.append(ch)
        prev = ch
    return unicodedata.normalize("NFC", "".join(out))


def normalize_title(text) -> str:
    """Canonical form: case-folded, punctuation-free, season labels unified.

    Examples
    --------
    >>> normalize_title("  The  Kapil Sharma Show! ")
    'kapil sharma show'
    >>> normalize_title("Bigg Boss Season 20")
    'bigg boss 20'
    """
    if text is None:
        return ""
    s = unicodedata.normalize("NFKC", str(text))
    s = _strip_latin_accents(s).casefold()
    s = s.translate(_APOSTROPHES)
    s = s.replace("&", " and ")
    s = re.sub(r"(?<=\d)\.(?=\d)", "\x00", s)  # keep decimals such as 1.5
    s = "".join(
        ch if (unicodedata.category(ch)[0] in "LNM" or ch == "\x00") else " " for ch in s
    )
    s = s.replace("\x00", ".")
    s = re.sub(r"\s+", " ", s).strip()
    s = _ORDINAL_SEASON.sub(r"\1", s)
    s = _SEASON_WORD.sub(lambda m: m.group(1).lstrip("0") or "0", s)
    s = re.sub(r"\s+", " ", s).strip()
    if s.startswith("the ") and len(s.split()) > 1:
        s = s[4:]
    return s


def compact(normalized: str) -> str:
    """Remove all spaces from an already-normalised string."""
    return normalized.replace(" ", "")


def compact_key(text) -> str:
    return compact(normalize_title(text))


def digit_tokens(normalized: str) -> Tuple[str, ...]:
    """All numbers in a normalised title, leading zeros removed."""
    out = []
    for d in _DIGITS.findall(normalized):
        if "." not in d:
            d = d.lstrip("0") or "0"
        out.append(d)
    return tuple(out)


def strip_trailing_number(normalized: str) -> str:
    return _TRAILING_NUMBER.sub("", normalized).strip()

# ======================================================================
# module: aliases
# ======================================================================
"""Recognised title variants (aliases) and forbidden pairs (exceptions).

Add new aliases here, or at run time via ``AliasTable.add`` / the JSON box in
the web app.  Each canonical title maps to a list of variants that should be
treated as the same show.  Numbers (season labels) are handled separately, so
list only the base title ("Bigg Boss", not "Bigg Boss 20").

JSON format accepted by ``load_custom_rules``::

    {
      "aliases":    {"Bigg Boss": ["BB", "Big Boss"]},
      "exceptions": [["Bigg Boss", "Bigg Boss OTT"]]
    }

Exceptions are pairs that must NEVER be matched, however similar they look.
"""


import json
from typing import Dict, Iterable, List, Tuple



DEFAULT_ALIASES: Dict[str, List[str]] = {
    "Bigg Boss": ["Big Boss", "Bigboss", "BB"],
    "The Kapil Sharma Show": ["TKSS"],
    "Kaun Banega Crorepati": ["KBC", "Kaun Banega Crorpati"],
    "Taarak Mehta Ka Ooltah Chashmah": [
        "TMKOC",
        "Tarak Mehta Ka Ulta Chashma",
        "Taarak Mehta",
        "Tarak Mehta",
    ],
    "Yeh Rishta Kya Kehlata Hai": ["YRKKH"],
    "Khatron Ke Khiladi": ["KKK", "Fear Factor Khatron Ke Khiladi"],
    "Game of Thrones": ["GOT"],
    "Money Heist": ["La Casa de Papel"],
    "MTV Roadies": ["Roadies"],
    "MTV Splitsvilla": ["Splitsvilla"],
    "Naagin": ["Nagin"],
    "Anupamaa": ["Anupama"],
}

# Pairs that must never match (checked before any other rule).
DEFAULT_EXCEPTIONS: List[Tuple[str, str]] = []


class AliasTable:
    """Maps compact title keys to a shared group id."""

    def __init__(self, groups: Dict[str, Iterable[str]] | None = None):
        self._key_to_group: Dict[str, str] = {}
        for canonical, variants in (groups or {}).items():
            self.add(canonical, variants)

    def add(self, canonical: str, variants: Iterable[str] = ()) -> None:
        base = strip_trailing_number(normalize_title(canonical))
        gid = compact(base)
        if not gid:
            return
        gid = self._key_to_group.get(gid, gid)
        for text in [canonical, *variants]:
            key = compact(strip_trailing_number(normalize_title(text)))
            if key:
                self._key_to_group[key] = gid

    def group_of(self, key: str) -> str:
        return self._key_to_group.get(key, key)

    def __len__(self) -> int:
        return len(self._key_to_group)


def default_alias_table() -> AliasTable:
    return AliasTable(DEFAULT_ALIASES)


def load_custom_rules(text: str) -> Tuple[Dict[str, List[str]], List[Tuple[str, str]]]:
    """Parse the optional JSON rules.  Raises ``ValueError`` with a clear message."""
    text = (text or "").strip()
    if not text:
        return {}, []
    try:
        data = json.loads(text)
    except json.JSONDecodeError as exc:
        raise ValueError(f"Custom rules are not valid JSON: {exc}") from exc
    if not isinstance(data, dict):
        raise ValueError("Custom rules must be a JSON object with 'aliases' and/or 'exceptions'.")
    aliases = data.get("aliases", {}) or {}
    exceptions = data.get("exceptions", []) or []
    if not isinstance(aliases, dict) or not all(
        isinstance(v, list) for v in aliases.values()
    ):
        raise ValueError("'aliases' must map a title to a list of variants.")
    if not isinstance(exceptions, list) or not all(
        isinstance(p, (list, tuple)) and len(p) == 2 for p in exceptions
    ):
        raise ValueError("'exceptions' must be a list of [titleA, titleB] pairs.")
    return (
        {str(k): [str(x) for x in v] for k, v in aliases.items()},
        [(str(a), str(b)) for a, b in exceptions],
    )

# ======================================================================
# module: pairs
# ======================================================================
"""Configurable question-pair checks.

Question numbers differ from survey to survey, so the checks are data, not code.
The default list is the original six checks; the web app lets the user type any pairs, e.g.

    Q4 -> Q9
    Q4, Q12 > Q9        (several sources for one destination are fine)
    Q5 → Q12

A question can be a source or a destination, never both (sources are never modified).
"""


import re
from typing import List, Sequence, Tuple



Check = Tuple[int, str, str]  # (check number, source question, destination question)

DEFAULT_CHECKS: List[Check] = [
    (1, "Q2", "Q3"),
    (2, "Q2", "Q7"),
    (3, "Q6", "Q7"),
    (4, "Q2", "Q15"),
    (5, "Q6", "Q15"),
    (6, "Q14", "Q15"),
]

_QID = re.compile(r"^\s*(?:q|question)?\s*\.?\s*0*(\d{1,3})\s*$", re.I)
_ARROW = re.compile(r"\s*(?:->|=>|-->|→|➜|>|to)\s*", re.I)


def normalize_qid(text: str) -> str:
    """'q2', 'Question 02', '2' -> 'Q2'.  Raises ValueError for anything else."""
    m = _QID.match(str(text))
    if not m:
        raise ValueError(f"'{text}' is not a question number (use e.g. Q2).")
    return f"Q{int(m.group(1))}"


def pairs_to_text(checks: Sequence[Check]) -> str:
    return "\n".join(f"{s} -> {d}" for _, s, d in checks)


def parse_pairs(text: str) -> List[Check]:
    """Parse one 'source -> destination' pair per line.

    The part before the arrow may list several sources separated by commas; each becomes its own
    check.  Duplicates are dropped, numbering follows the order typed.
    """
    errors: List[str] = []
    out: List[Check] = []
    seen = set()
    for lineno, raw in enumerate((text or "").splitlines(), start=1):
        line = raw.split("#")[0].strip()
        if not line:
            continue
        parts = _ARROW.split(line)
        if len(parts) != 2 or not parts[0].strip() or not parts[1].strip():
            errors.append(f"Line {lineno} ('{raw.strip()}'): write it as  Q2 -> Q3")
            continue
        try:
            dst = normalize_qid(parts[1])
            sources = [normalize_qid(s) for s in re.split(r"[,&+/]|\band\b", parts[0]) if s.strip()]
        except ValueError as exc:
            errors.append(f"Line {lineno}: {exc}")
            continue
        for src in sources:
            if src == dst:
                errors.append(f"Line {lineno}: {src} cannot be compared with itself.")
            elif (src, dst) not in seen:
                seen.add((src, dst))
                out.append((len(out) + 1, src, dst))
    if not out and not errors:
        errors.append("Enter at least one question pair, e.g.  Q2 -> Q3")
    check_pairs(out, errors)
    if errors:
        raise MappingError(errors)
    return out


def check_pairs(checks: Sequence[Check], errors: List[str]) -> None:
    both = sorted({s for _, s, _ in checks} & {d for _, _, d in checks}, key=lambda q: int(q[1:]))
    for q in both:
        errors.append(
            f"{q} is used both as a source and as a destination. Source answers are never modified, "
            "so a question must be only one of the two."
        )


def sources_of(checks: Sequence[Check]) -> List[str]:
    return _ordered(s for _, s, _ in checks)


def destinations_of(checks: Sequence[Check]) -> List[str]:
    return _ordered(d for _, _, d in checks)


def questions_of(checks: Sequence[Check]) -> List[str]:
    """Every question used, sorted by question number."""
    return sorted({q for _, s, d in checks for q in (s, d)}, key=lambda q: int(q[1:]))


def _ordered(items) -> List[str]:
    seen, out = set(), []
    for q in items:
        if q not in seen:
            seen.add(q)
            out.append(q)
    return sorted(out, key=lambda q: int(q[1:]))

# ======================================================================
# module: matching
# ======================================================================
"""Title-matching engine.

Matching order (first rule that fires wins):

1. ``exact``       - equal after trimming and case-folding.
2. ``normalized``  - equal after punctuation/space/season-label/"The" normalisation,
                     including joined vs. separated words.
3. ``alias``       - both titles belong to the same recognised variant group.
4. ``fuzzy``       - conservative similarity >= threshold (default 92).

Guards against false positives
------------------------------
* Different numbers never fuzzy-match ("Bigg Boss 19" vs "Bigg Boss 20").
* A number on one side only ("Bigg Boss" vs "Bigg Boss 20") is sent to review.
* Short titles (< ``min_fuzzy_len`` letters) need a much higher score, and very
  short titles (< ``short_fuzzy_len``) are never fuzzy-matched at all.
* If several destination headers qualify equally well the title is *ambiguous*
  and goes to the manual-review list instead of being applied.
"""


import itertools
from dataclasses import dataclass, field
from typing import Dict, Iterable, List, Optional, Sequence, Tuple









try:  # RapidFuzz is the preferred scorer (used on Streamlit Cloud).
    from rapidfuzz import fuzz as _rf_fuzz
except Exception:  # pragma: no cover - fallback keeps the engine usable anywhere
    _rf_fuzz = None

# Match types ---------------------------------------------------------------
EXACT = "exact"
NORMALIZED = "normalized"
ALIAS = "alias"
FUZZY = "fuzzy"
REVIEW = "review"
NONE = "none"

AUTO_TYPES = (EXACT, NORMALIZED, ALIAS, FUZZY)
# Higher number = stronger evidence.  Used to resolve several titles that
# point at the same destination cell (exact > normalized > alias > fuzzy).
PRIORITY: Dict[str, int] = {EXACT: 4, NORMALIZED: 3, ALIAS: 2, FUZZY: 1}
# Yellow for exact; green for every other automatic match type.
COLOR_FOR: Dict[str, str] = {EXACT: "Yellow", NORMALIZED: "Green", ALIAS: "Green", FUZZY: "Green"}

DEFAULT_THRESHOLD = 92.0


# Similarity ----------------------------------------------------------------
def _lcs_len(a: str, b: str) -> int:
    if len(a) < len(b):
        a, b = b, a
    prev = [0] * (len(b) + 1)
    for ca in a:
        cur = [0]
        for j, cb in enumerate(b, 1):
            cur.append(prev[j - 1] + 1 if ca == cb else max(prev[j], cur[j - 1]))
        prev = cur
    return prev[-1]


def similarity(a: str, b: str) -> float:
    """Normalised Indel similarity 0-100 (identical to ``rapidfuzz.fuzz.ratio``)."""
    if not a and not b:
        return 100.0
    if not a or not b:
        return 0.0
    if _rf_fuzz is not None:
        return float(_rf_fuzz.ratio(a, b))
    return 200.0 * _lcs_len(a, b) / (len(a) + len(b))


def scorer_name() -> str:
    return "rapidfuzz" if _rf_fuzz is not None else "built-in fallback (same formula)"


# Result types ----------------------------------------------------------------
@dataclass(frozen=True)
class MatchResult:
    type: str
    score: float = 0.0
    reason: str = ""


NO_MATCH = MatchResult(NONE, 0.0, "")


@dataclass(frozen=True)
class Candidate:
    col: int
    header: str
    result: MatchResult


@dataclass
class Resolution:
    """Outcome of comparing one source title with every header of a question."""

    status: str  # 'matched' | 'ambiguous' | 'review' | 'none'
    candidates: List[Candidate] = field(default_factory=list)
    reason: str = ""

    @property
    def best(self) -> Optional[Candidate]:
        return self.candidates[0] if self.status == "matched" and self.candidates else None


_header_ids = itertools.count(1)


class HeaderSet:
    """Destination-question headers: a list of (column index, header text)."""

    def __init__(self, items: Iterable[Tuple[int, str]]):
        self.items: List[Tuple[int, str]] = list(items)
        self.uid = next(_header_ids)

    def __iter__(self):
        return iter(self.items)

    def __len__(self):
        return len(self.items)


class TitleMatcher:
    def __init__(
        self,
        threshold: float = DEFAULT_THRESHOLD,
        aliases: AliasTable | None = None,
        exceptions: Iterable[Tuple[str, str]] = (),
        min_fuzzy_len: int = 7,
        short_fuzzy_len: int = 5,
        short_threshold: float = 96.0,
        tie_margin: float = 2.0,
        review_margin: float = 12.0,
    ):
        if not 50.0 <= float(threshold) <= 100.0:
            raise ValueError("Fuzzy threshold must be between 50 and 100.")
        self.threshold = float(threshold)
        self.aliases = aliases if aliases is not None else default_alias_table()
        self.min_fuzzy_len = min_fuzzy_len
        self.short_fuzzy_len = short_fuzzy_len
        self.short_threshold = max(short_threshold, self.threshold)
        self.tie_margin = tie_margin
        self.review_floor = max(60.0, self.threshold - review_margin)
        self._exceptions: set = set()
        for a, b in exceptions:
            self.add_exception(a, b)
        self._cache: Dict[Tuple[str, int], Resolution] = {}
        self._norm_cache: Dict[str, Tuple[str, str, Tuple[str, ...], str]] = {}

    # -- extensibility ------------------------------------------------------
    def add_alias(self, canonical: str, variants: Iterable[str] = ()) -> None:
        self.aliases.add(canonical, variants)
        self._cache.clear()

    def add_exception(self, a: str, b: str) -> None:
        ka, kb = compact(normalize_title(a)), compact(normalize_title(b))
        if ka and kb:
            self._exceptions.add(frozenset((ka, kb)))
        self._cache.clear()

    # -- helpers --------------------------------------------------------------
    def _forms(self, text: str) -> Tuple[str, str, Tuple[str, ...], str]:
        """(normalised, compact, digits, alias-group-of-base)."""
        hit = self._norm_cache.get(text)
        if hit is None:
            norm = normalize_title(text)
            base = compact(strip_trailing_number(norm))
            hit = (norm, compact(norm), digit_tokens(norm), self.aliases.group_of(base) if base else "")
            self._norm_cache[text] = hit
        return hit

    # -- pairwise comparison --------------------------------------------------
    def compare(self, source: str, header: str) -> MatchResult:
        if source is None or header is None:
            return NO_MATCH
        s_raw, h_raw = str(source).strip(), str(header).strip()
        if not s_raw or not h_raw:
            return NO_MATCH

        s_norm, s_key, s_dig, s_grp = self._forms(s_raw)
        h_norm, h_key, h_dig, h_grp = self._forms(h_raw)
        if not s_key or not h_key:
            return NO_MATCH

        if frozenset((s_key, h_key)) in self._exceptions:
            return NO_MATCH

        if s_raw.casefold() == h_raw.casefold():
            return MatchResult(EXACT, 100.0, "identical ignoring case and outer spaces")

        if s_key == h_key:
            return MatchResult(NORMALIZED, 100.0, "identical after normalisation")

        if s_grp and s_grp == h_grp and s_dig == h_dig:
            return MatchResult(ALIAS, 100.0, "recognised title variant")

        # ---- numbers: never fuzzy-match different seasons ------------------
        if s_dig != h_dig:
            if s_dig and h_dig:
                return NO_MATCH  # e.g. Bigg Boss 19 vs Bigg Boss 20: different titles
            s_base = compact(strip_trailing_number(s_norm))
            h_base = compact(strip_trailing_number(h_norm))
            same_base = bool(s_base) and (
                s_base == h_base
                or (s_grp and s_grp == h_grp)
                or similarity(s_base, h_base) >= self.threshold
            )
            if same_base:
                return MatchResult(
                    REVIEW,
                    100.0 if s_base == h_base else round(similarity(s_base, h_base), 1),
                    "season/number appears on one side only - confirm it is the same title",
                )
            return NO_MATCH

        # ---- fuzzy -----------------------------------------------------------
        score = round(similarity(s_key, h_key), 2)
        shortest = min(len(s_key), len(h_key))
        if shortest < self.short_fuzzy_len:
            if score >= self.review_floor:
                return MatchResult(
                    REVIEW, score, f"very short title (<{self.short_fuzzy_len} letters): never auto-matched"
                )
            return NO_MATCH
        required = self.threshold if shortest >= self.min_fuzzy_len else self.short_threshold
        if score >= required:
            return MatchResult(FUZZY, score, f"similarity {score:.1f}% >= {required:.0f}%")
        if score >= self.review_floor:
            return MatchResult(
                REVIEW, score, f"similar ({score:.1f}%) but below the required {required:.0f}%"
            )
        return NO_MATCH

    # -- one source title vs. all destination headers -------------------------
    def resolve(self, title: str, headers: HeaderSet) -> Resolution:
        key = (str(title).strip(), headers.uid)
        hit = self._cache.get(key)
        if hit is not None:
            return hit
        out = self._resolve(key[0], headers)
        self._cache[key] = out
        return out

    def _resolve(self, title: str, headers: HeaderSet) -> Resolution:
        cands: List[Candidate] = []
        for col, text in headers:
            res = self.compare(title, text)
            if res.type != NONE:
                cands.append(Candidate(col, text, res))
        if not cands:
            return Resolution("none")

        qualified = [c for c in cands if c.result.type in AUTO_TYPES]
        if qualified:
            top_priority = max(PRIORITY[c.result.type] for c in qualified)
            top = [c for c in qualified if PRIORITY[c.result.type] == top_priority]
            if top_priority == PRIORITY[FUZZY]:
                top.sort(key=lambda c: (-c.result.score, c.col))
                top = [c for c in top if top[0].result.score - c.result.score <= self.tie_margin]
            else:
                top.sort(key=lambda c: c.col)
            if len(top) == 1:
                return Resolution("matched", top)
            return Resolution(
                "ambiguous",
                top,
                f"{len(top)} destination headers match equally well - choose manually",
            )

        near = sorted(
            (c for c in cands if c.result.type == REVIEW),
            key=lambda c: (-c.result.score, c.col),
        )[:5]
        if near:
            return Resolution("review", near, near[0].result.reason)
        return Resolution("none")

# ======================================================================
# module: workbook_io
# ======================================================================
"""Reading workbooks into lightweight snapshots + unsupported-feature warnings."""


import io
import re
import zipfile
from dataclasses import dataclass, field
from typing import Dict, List, Optional, Tuple

from openpyxl import load_workbook



Merged = Tuple[int, int, int, int]  # min_row, min_col, max_row, max_col


@dataclass
class Snapshot:
    """In-memory copy of one worksheet (1-based accessors).

    ``values``   - what a person sees (cached formula results are used for titles).
    ``raw``      - what is stored in the file (formulas appear as '=...' strings);
                   destination cells are classified from this so formulas are never overwritten.
    """

    sheet: str
    values: List[list]
    raw: List[list]
    max_row: int
    max_col: int
    merged: List[Merged] = field(default_factory=list)

    def value(self, row: int, col: int):
        if 1 <= row <= self.max_row and 1 <= col <= self.max_col:
            return self.values[row - 1][col - 1]
        return None

    def raw_value(self, row: int, col: int):
        if 1 <= row <= self.max_row and 1 <= col <= self.max_col:
            return self.raw[row - 1][col - 1]
        return None

    def merged_at(self, row: int, col: int) -> Optional[Merged]:
        for m in self.merged:
            if m[0] <= row <= m[2] and m[1] <= col <= m[3]:
                return m
        return None


def is_blank(v) -> bool:
    return v is None or (isinstance(v, str) and v.strip() == "")


def _open(data: bytes, data_only: bool, read_only: bool = False):
    if not data or data[:2] != b"PK":
        raise WorkbookError(
            "This does not look like an .xlsx workbook. Please upload an Excel .xlsx file "
            "(older .xls files must be re-saved as .xlsx first)."
        )
    try:
        return load_workbook(io.BytesIO(data), data_only=data_only, read_only=read_only)
    except zipfile.BadZipFile as exc:
        raise WorkbookError("The file is corrupted or is not a valid .xlsx workbook.") from exc
    except Exception as exc:  # openpyxl raises many different errors
        raise WorkbookError(f"The workbook could not be read: {exc}") from exc


def open_workbook(data: bytes, data_only: bool = False, read_only: bool = False):
    return _open(data, data_only, read_only)


def sheet_names(data: bytes) -> List[str]:
    return list(_open(data, True).sheetnames)


def _grid(ws) -> List[list]:
    return [
        list(r)
        for r in ws.iter_rows(
            min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column, values_only=True
        )
    ]


def _has_formula(grid: List[list]) -> bool:
    return any(isinstance(v, str) and v.startswith("=") for row in grid for v in row)


def snapshot_from_sheet(ws_e, data: bytes) -> Snapshot:
    """Build a snapshot from an already-loaded (formula) worksheet.

    ``raw`` holds what is stored; ``values`` equals ``raw`` except that formula cells show
    their cached results.  The cached-values workbook is only opened when the sheet
    really contains formulas, so ordinary survey exports are loaded once.
    """
    raw = _grid(ws_e)
    values = [row[:] for row in raw]
    if _has_formula(raw):
        try:
            wb_v = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
            cached = [list(r) for r in wb_v[ws_e.title].iter_rows(values_only=True)]
            wb_v.close()
        except Exception:  # cached values are a convenience only
            cached = []
        for r, row in enumerate(values):
            for c, v in enumerate(row):
                if isinstance(v, str) and v.startswith("="):
                    got = cached[r][c] if r < len(cached) and c < len(cached[r]) else None
                    row[c] = got
    merged = [
        (m.min_row, m.min_col, m.max_row, m.max_col) for m in ws_e.merged_cells.ranges
    ]
    return Snapshot(ws_e.title, values, raw, ws_e.max_row, ws_e.max_column, merged)


def load_snapshot(data: bytes, sheet: str) -> Snapshot:
    wb = _open(data, False)
    if sheet not in wb.sheetnames:
        raise WorkbookError(f"Worksheet '{sheet}' was not found in the workbook.")
    return snapshot_from_sheet(wb[sheet], data)


def load_all_snapshots(data: bytes) -> Dict[str, Snapshot]:
    """Snapshots of every sheet (one workbook load), for inspection."""
    wb = _open(data, False)
    return {ws.title: snapshot_from_sheet(ws, data) for ws in wb.worksheets}


# --------------------------------------------------------------------------
# Unsupported-feature detection
# --------------------------------------------------------------------------
_FEATURES = [
    ("xl/charts/", "Charts are present. OpenPyXL cannot read charts, so they may be LOST in the corrected workbook."),
    ("xl/drawings/", "Drawings/shapes/images are present and may be lost or altered in the corrected workbook."),
    ("xl/media/", "Embedded images are present; they may be lost unless the Pillow library is available."),
    ("xl/vbaProject.bin", "A VBA macro project is present (macros are not kept in .xlsx output)."),
    ("xl/slicers/", "Slicers are present and may be lost."),
    ("xl/timelines/", "Timelines are present and may be lost."),
    ("xl/ctrlProps/", "Form controls are present and may be lost."),
    ("xl/embeddings/", "Embedded objects are present and may be lost."),
    ("xl/threadedComments/", "Threaded comments are present; they may be converted to plain notes or lost."),
    ("xl/pivotTables/", "Pivot tables are present; they are kept but should be checked after processing."),
]


def detect_unsupported_features(data: bytes) -> List[str]:
    """Return warnings about workbook parts that OpenPyXL may not round-trip."""
    warnings: List[str] = []
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            names = zf.namelist()
            for prefix, message in _FEATURES:
                if any(n.startswith(prefix) or n == prefix for n in names):
                    warnings.append(message)
            ext_hits = 0
            sparkline = False
            for n in names:
                if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", n):
                    xml = zf.read(n)
                    if b"<extLst" in xml:
                        ext_hits += 1
                    if b"sparkline" in xml:
                        sparkline = True
            if sparkline:
                warnings.append("Sparklines are present and will be lost.")
            if ext_hits:
                warnings.append(
                    "Some worksheets use Excel extension features (e.g. newer conditional formatting "
                    "or data-validation options) that may be dropped."
                )
    except zipfile.BadZipFile:
        pass
    return warnings

# ======================================================================
# module: inspector
# ======================================================================
"""Automatic detection of question layout + validation of (possibly edited) mappings.

Nothing here hard-codes column letters, header rows or sheet names.  Detection is
heuristic, so every result carries ``confident`` and ``issues``; the web app
always requires the user to confirm the mapping before anything is changed.
"""


import re
from dataclasses import dataclass, field
from typing import Dict, List, Optional, Sequence, Tuple













# Defaults (the original six checks).  Real runs pass their own ``checks`` list.
SOURCE_QUESTIONS = tuple(sources_of(DEFAULT_CHECKS))
DEST_QUESTIONS = tuple(destinations_of(DEFAULT_CHECKS))
ALL_QUESTIONS = tuple(questions_of(DEFAULT_CHECKS))

SCAN_ROWS = 15  # how many top rows are searched for question labels
HEADER_SEARCH_ROWS = 5
DATA_SAMPLE_ROWS = 200

QUESTION_RE = re.compile(r"^\s*(?:q|question)\s*\.?\s*0*(\d{1,3})(?!\d)", re.I)


def question_number(value) -> Optional[int]:
    """'Q2', 'Q2_1', 'Q 2.', 'Question 2 - Which ...' -> 2."""
    if not isinstance(value, str):
        return None
    m = QUESTION_RE.match(value)
    return int(m.group(1)) if m else None


@dataclass
class QuestionMapping:
    qid: str
    columns: List[int] = field(default_factory=list)
    label_cell: str = ""
    merged: bool = False
    note: str = ""

    @property
    def spec(self) -> str:
        return columns_to_spec(self.columns)


@dataclass
class SheetLayout:
    sheet: str
    label_row: int = 1
    header_row: int = 2
    first_row: int = 3
    last_row: int = 3
    questions: Dict[str, QuestionMapping] = field(default_factory=dict)
    confident: bool = False
    issues: List[str] = field(default_factory=list)

    def cols(self, qid: str) -> List[int]:
        q = self.questions.get(qid)
        return list(q.columns) if q else []

    @property
    def found(self) -> int:
        return sum(1 for q in self.questions if self.cols(q))


def build_layout(
    sheet: str,
    label_row: int,
    header_row: int,
    first_row: int,
    last_row: int,
    specs: Dict[str, str],
) -> SheetLayout:
    """Create a layout from user input; column specs use Excel notation ('BK:BQ')."""
    layout = SheetLayout(sheet, int(label_row), int(header_row), int(first_row), int(last_row))
    errors = []
    for raw_qid, spec in specs.items():
        qid = normalize_qid(raw_qid)
        try:
            cols = parse_columns(spec)
        except ValueError as exc:
            errors.append(f"{qid}: {exc}")
            cols = []
        layout.questions[qid] = QuestionMapping(qid, cols, note="user-confirmed")
    if errors:
        raise MappingError(errors)
    layout.confident = True
    return layout


# --------------------------------------------------------------------------
# Detection
# --------------------------------------------------------------------------
def _nonempty_in_column(snap: Snapshot, col: int, r0: int, r1: int) -> bool:
    return any(not is_blank(snap.value(r, col)) for r in range(r0, r1 + 1))


def _looks_like_header(snap: Snapshot, row: int, cols: Sequence[int]) -> bool:
    """True when the row holds text titles instead of 0/1/2 codes."""
    vals = [snap.value(row, c) for c in cols]
    vals = [v for v in vals if not is_blank(v)]
    if len(vals) < 2:  # a single odd value is data, not a sub-header row
        return False
    texty = 0
    for v in vals:
        if isinstance(v, str) and v.strip() not in {"0", "1", "2"}:
            try:
                float(v)
            except ValueError:
                texty += 1
    return texty / len(vals) >= 0.5


def inspect_snapshot(snap: Snapshot, checks=None) -> SheetLayout:
    """Detect the layout.  Every question label found (Q1, Q2, ... any number) is mapped;
    the questions named in ``checks`` (default: the six standard ones) are the *required* ones."""
    checks = list(checks if checks is not None else DEFAULT_CHECKS)
    required = questions_of(checks)
    dest_qids = destinations_of(checks)
    layout = SheetLayout(snap.sheet)
    for q in required:
        layout.questions[q] = QuestionMapping(q)

    # 1. find the label row ---------------------------------------------------
    best_row, best_count, row_labels = None, 0, {}
    for r in range(1, min(snap.max_row, SCAN_ROWS) + 1):
        labels = []
        for c in range(1, snap.max_col + 1):
            qn = question_number(snap.value(r, c))
            if qn is not None:
                labels.append((c, qn))
        row_labels[r] = labels
        count = len({qn for _, qn in labels})
        if count > best_count:
            best_row, best_count = r, count

    if best_row is None:
        layout.issues.append(
            "No question labels (Q2, Q3, Q6, Q7, Q14, Q15) were found in the top "
            f"{SCAN_ROWS} rows. Please enter the column ranges manually."
        )
        layout.last_row = snap.max_row
        return layout

    layout.label_row = best_row
    labels = sorted(row_labels[best_row])

    # 2. group consecutive labels of the same question into column segments ----
    groups: List[Tuple[int, int, int]] = []  # (qnum, start_col, label_count)
    for col, qn in labels:
        if groups and groups[-1][0] == qn:
            continue
        groups.append((qn, col, 0))

    found_before: Dict[str, int] = {}
    for i, (qn, start, _) in enumerate(groups):
        end = groups[i + 1][1] - 1 if i + 1 < len(groups) else snap.max_col
        merged = snap.merged_at(best_row, start)
        is_merged = bool(merged and merged[3] > merged[1])
        if is_merged:
            end = max(min(end, snap.max_col), merged[3])
        # trim trailing columns with no content at all
        r_hi = min(snap.max_row, best_row + DATA_SAMPLE_ROWS)
        while end > start and not _nonempty_in_column(snap, end, best_row, r_hi):
            end -= 1
        qid = f"Q{qn}"
        if qid in found_before:
            layout.issues.append(
                f"{qid} appears more than once in row {best_row}; the first occurrence "
                f"({col_letter(found_before[qid])}) was used - please confirm."
            )
            continue
        found_before[qid] = start
        layout.questions[qid] = QuestionMapping(
            qid,
            list(range(start, end + 1)),
            label_cell=addr(best_row, start),
            merged=is_merged,
            note="merged label cell" if is_merged else "",
        )

    missing = [q for q in required if not layout.cols(q)]
    if missing:
        layout.issues.append(
            f"Could not find label(s) for: {', '.join(missing)}. Enter those column ranges manually "
            "(or correct the question numbers in the question-pair list)."
        )

    # 3. header row (titles of the destination questions) ----------------------
    dest_cols = [c for q in dest_qids for c in layout.cols(q)]
    header_row = best_row + 1
    if dest_cols:
        chosen, best_guess, best_share = None, None, 0.0
        for r in range(best_row + 1, best_row + 1 + HEADER_SEARCH_ROWS):
            if r > snap.max_row:
                break
            filled = sum(
                1 for c in dest_cols if isinstance(snap.value(r, c), str) and snap.value(r, c).strip()
            )
            share = filled / len(dest_cols)
            if share >= 0.6:
                chosen = r
                break
            if share > best_share:
                best_share, best_guess = share, r
        if chosen is not None:
            header_row = chosen
        elif best_guess is not None:
            header_row = best_guess
            layout.issues.append(
                f"Title headers in row {best_guess} are only partly filled; please check the header row."
            )
        else:
            layout.issues.append("No row below the question labels contains title headers.")
    else:
        layout.issues.append("Header row could not be determined because no destination question was found.")
    layout.header_row = header_row

    # 4. first respondent row ----------------------------------------------------
    first = header_row + 1
    skipped = 0
    while skipped < 3 and first <= snap.max_row and dest_cols and _looks_like_header(snap, first, dest_cols):
        first += 1
        skipped += 1
    if skipped:
        layout.issues.append(
            f"Skipped {skipped} text-only row(s) below the header row; respondent data is assumed to "
            f"start at row {first} - please confirm."
        )
    layout.first_row = first

    # 5. last respondent row -------------------------------------------------------
    mapped = [c for q in layout.questions for c in layout.cols(q)]
    last = snap.max_row
    while last > first and not any(not is_blank(snap.value(last, c)) for c in mapped):
        last -= 1
    layout.last_row = max(last, first)

    layout.confident = not layout.issues and all(layout.cols(q) for q in required)
    return layout


def inspect_snapshots(snaps: Dict[str, Snapshot], checks=None) -> Dict[str, SheetLayout]:
    return {name: inspect_snapshot(s, checks) for name, s in snaps.items()}


def best_sheet(layouts: Dict[str, SheetLayout]) -> str:
    """Sheet with the most detected questions (confident sheets win ties)."""
    return max(layouts, key=lambda n: (layouts[n].found, layouts[n].confident, -list(layouts).index(n)))


# --------------------------------------------------------------------------
# Validation
# --------------------------------------------------------------------------
def dest_headers(snap: Snapshot, layout: SheetLayout, qid: str) -> List[Tuple[int, str]]:
    """(column, title) for every non-empty header cell of a destination question."""
    out = []
    for c in layout.cols(qid):
        v = snap.value(layout.header_row, c)
        if v is None or isinstance(v, bool):
            continue
        if isinstance(v, float) and v.is_integer():
            v = int(v)
        text = str(v).strip()
        if text:
            out.append((c, text))
    return out


def validate_layout(snap: Snapshot, layout: SheetLayout, checks=None) -> Tuple[List[str], List[str]]:
    """Return (errors, warnings).  Errors block processing.  ``checks`` = the question pairs in use."""
    checks = list(checks if checks is not None else DEFAULT_CHECKS)
    used = questions_of(checks)
    dest_qids = destinations_of(checks)
    errors: List[str] = []
    warnings: List[str] = []
    check_pairs(checks, errors)

    for qid in used:
        if not layout.cols(qid):
            errors.append(f"{qid} has no columns mapped. Enter its column range (e.g. BK:BQ).")
    if layout.header_row < 1:
        errors.append("Header row must be 1 or greater.")
    if layout.first_row <= layout.header_row:
        errors.append("The first respondent row must be below the title-header row.")
    if layout.last_row < layout.first_row:
        errors.append("The last respondent row must not be above the first respondent row.")
    if layout.first_row > snap.max_row:
        errors.append(
            f"The first respondent row ({layout.first_row}) is beyond the end of the sheet "
            f"(last used row is {snap.max_row})."
        )

    # a column may belong to one question only -> source answers can never be edited
    owner: Dict[int, str] = {}
    for qid in used:
        for c in layout.cols(qid):
            if c in owner and owner[c] != qid:
                errors.append(
                    f"Column {col_letter(c)} is mapped to both {owner[c]} and {qid}. "
                    "Question ranges must not overlap (source answers must never be modified)."
                )
            owner.setdefault(c, qid)

    for qid in dest_qids:
        if layout.cols(qid) and not dest_headers(snap, layout, qid):
            errors.append(
                f"{qid}: no title headers were found in row {layout.header_row}. "
                "Check the header row number or the column range."
            )
        else:
            titles = [t.casefold() for _, t in dest_headers(snap, layout, qid)]
            dupes = sorted({t for t in titles if titles.count(t) > 1})
            if dupes:
                warnings.append(
                    f"{qid} has duplicate title headers ({', '.join(dupes[:5])}); "
                    "matches to them will be treated as ambiguous and sent to review."
                )
            empties = len(layout.cols(qid)) - len(titles)
            if empties:
                warnings.append(f"{qid}: {empties} column(s) have an empty header and will be ignored.")

    for qid in used:
        beyond = [c for c in layout.cols(qid) if c > snap.max_col]
        if beyond:
            warnings.append(f"{qid}: column(s) beyond the last used column are mapped (they are empty).")
    return errors, warnings

# ======================================================================
# module: processor
# ======================================================================
"""Respondent-level matching and correction of destination cells.

Pipeline
--------
1. ``analyze``   - for each configured check (default: the six standard pairs), read every respondent's source titles
                   and resolve them against the destination headers.
2. ``decide``    - group matches by destination cell (across all checks), pick one
                   deterministic winner per cell (exact > normalized > alias > fuzzy,
                   then higher score, then lower check number) and apply the cell rules.
3. ``apply``     - write 0 + fill colour into a *copy* of the workbook (OpenPyXL).
4. ``verify``    - re-open the saved bytes and prove that only the intended cells differ.

The uploaded bytes are never modified; a new workbook is produced.
"""


import gc
import hashlib
import io
import re
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import datetime
from typing import Dict, List, Optional, Tuple

from openpyxl.styles import PatternFill































# (check number, source question, destination question) - exactly the six required checks
CHECKS: List[Tuple[int, str, str]] = list(DEFAULT_CHECKS)  # default; callers may pass their own

YELLOW_HEX = "FFFF00"
GREEN_HEX = "00B050"
COLOR_HEX = {"Yellow": YELLOW_HEX, "Green": GREEN_HEX}

# Several titles in one source cell may be separated by these characters.
SPLIT_RE = re.compile(r"[\n\r;|]+")

PRIORITY_RULE_TEXT = (
    "When several source titles (or several checks) point at the same destination cell, the cell is "
    "changed and counted once. The winning match is chosen deterministically: exact beats normalized, "
    "normalized beats recognised-variant (alias), alias beats fuzzy; ties go to the higher similarity "
    "score, then the lower check number, then the leftmost source column. The cell colour follows the "
    "winning match (yellow = exact, green = anything else). Other matches to that cell are logged as "
    "'duplicate' and not counted as changes."
)


# --------------------------------------------------------------------------
# Value classification
# --------------------------------------------------------------------------
ONE, TWO, ZERO, BLANK, FORMULA, OTHER = "one", "two", "zero", "blank", "formula", "other"


def classify_value(v) -> str:
    """Classify a destination cell value according to the correction rules."""
    if v is None:
        return BLANK
    if isinstance(v, bool):
        return OTHER
    if isinstance(v, (int, float)):
        if v == 1:
            return ONE
        if v == 2:
            return TWO
        if v == 0:
            return ZERO
        return OTHER
    if isinstance(v, str):
        s = v.strip()
        if s == "":
            return BLANK
        if s.startswith("="):
            return FORMULA
        return {"1": ONE, "2": TWO, "0": ZERO}.get(s, OTHER)
    return OTHER


# --------------------------------------------------------------------------
# Data classes
# --------------------------------------------------------------------------
@dataclass
class MatchEvent:
    check_no: int
    src_q: str
    dst_q: str
    row: int
    src_col: int
    src_title: str
    dst_col: int
    dst_header: str
    type: str
    score: float
    reason: str
    seq: int = 0


@dataclass
class ReviewEvent:
    check_no: int
    src_q: str
    dst_q: str
    row: int
    src_col: int
    src_title: str
    status: str  # 'ambiguous' | 'review'
    candidates: List[Candidate]
    reason: str


@dataclass
class Analysis:
    events: List[MatchEvent] = field(default_factory=list)
    reviews: List[ReviewEvent] = field(default_factory=list)
    titles_checked: Counter = field(default_factory=Counter)
    header_cells: Dict[str, Dict[int, str]] = field(default_factory=dict)


@dataclass
class Decision:
    row: int
    col: int
    winner: MatchEvent
    others: List[MatchEvent]
    original: object
    kind: str
    action: str
    final: object
    color: str  # '' when not changed


@dataclass
class ProcessResult:
    filename: str
    sheet: str
    processed_at: datetime
    threshold: float
    layout: SheetLayout
    corrected_bytes: bytes
    audit_rows: List[dict]
    review_rows: List[dict]
    per_check: List[dict]
    totals: dict
    changed_cells: List[dict]
    feature_warnings: List[str]
    layout_warnings: List[str]
    verification_issues: List[str]
    verified: bool
    scorer: str = ""
    checks: List[Tuple[int, str, str]] = field(default_factory=lambda: list(CHECKS))


# --------------------------------------------------------------------------
# Step 1: analysis
# --------------------------------------------------------------------------
def split_titles(value: str) -> List[str]:
    return [p.strip() for p in SPLIT_RE.split(value) if p.strip()]


def analyze(snap: Snapshot, layout: SheetLayout, matcher: TitleMatcher, checks=None) -> Analysis:
    checks = list(checks if checks is not None else CHECKS)
    out = Analysis()
    header_sets: Dict[str, HeaderSet] = {}
    for qid in {dq for _, _, dq in checks}:
        items = dest_headers(snap, layout, qid)
        header_sets[qid] = HeaderSet(items)
        out.header_cells[qid] = {c: t for c, t in items}

    seq = 0
    for check_no, sq, dq in checks:
        hs = header_sets[dq]
        src_cols = layout.cols(sq)
        for row in range(layout.first_row, layout.last_row + 1):
            for sc in src_cols:
                value = snap.value(row, sc)
                if not isinstance(value, str) or not value.strip():
                    continue
                for title in split_titles(value):
                    out.titles_checked[check_no] += 1
                    res = matcher.resolve(title, hs)
                    if res.status == "matched":
                        cand = res.best
                        seq += 1
                        out.events.append(
                            MatchEvent(
                                check_no, sq, dq, row, sc, title, cand.col, cand.header,
                                cand.result.type, cand.result.score, cand.result.reason, seq,
                            )
                        )
                    elif res.status in ("ambiguous", "review"):
                        out.reviews.append(
                            ReviewEvent(
                                check_no, sq, dq, row, sc, title, res.status,
                                list(res.candidates), res.reason,
                            )
                        )
    return out


# --------------------------------------------------------------------------
# Step 2: decisions
# --------------------------------------------------------------------------
def _winner_key(ev: MatchEvent):
    return (-PRIORITY[ev.type], -ev.score, ev.check_no, ev.src_col, ev.seq)


def decide(snap: Snapshot, analysis: Analysis) -> List[Decision]:
    groups: Dict[Tuple[int, int], List[MatchEvent]] = defaultdict(list)
    for ev in analysis.events:
        groups[(ev.row, ev.dst_col)].append(ev)

    decisions: List[Decision] = []
    for (row, col), evs in sorted(groups.items()):
        evs.sort(key=_winner_key)
        winner, others = evs[0], evs[1:]
        original = snap.raw_value(row, col)
        kind = classify_value(original)
        color = ""
        final = original
        if kind == ONE:
            final = "0" if isinstance(original, str) else 0
            color = COLOR_FOR[winner.type]
            action = "Changed 1 to 0"
        elif kind == TWO:
            action = "Preserved (value 2)"
        elif kind == ZERO:
            action = "No change (already 0)"
        elif kind == BLANK:
            action = "Skipped (blank cell)"
        elif kind == FORMULA:
            action = "Skipped (cell contains a formula)"
        else:
            action = "Skipped (other value)"
        decisions.append(Decision(row, col, winner, others, original, kind, action, final, color))
    return decisions


# --------------------------------------------------------------------------
# Reporting rows / statistics
# --------------------------------------------------------------------------
def _display(v):
    return "" if v is None else v


def build_audit_rows(layout: SheetLayout, decisions: List[Decision]) -> List[dict]:
    rows: List[dict] = []
    for d in decisions:
        for ev in [d.winner, *d.others]:
            is_winner = ev is d.winner
            if is_winner:
                action = d.action
                modified = d.kind == ONE
                color = d.color
            else:
                action = f"Duplicate match - cell already handled by Check {d.winner.check_no}"
                modified = False
                color = ""
            rows.append(
                {
                    "Worksheet": layout.sheet,
                    "Check": ev.check_no,
                    "Source Question": ev.src_q,
                    "Destination Question": ev.dst_q,
                    "Respondent Row": ev.row,
                    "Source Title": ev.src_title,
                    "Source Cell": addr(ev.row, ev.src_col),
                    "Destination Title Header": ev.dst_header,
                    "Destination Header Cell": addr(layout.header_row, ev.dst_col),
                    "Destination Cell": addr(d.row, d.col),
                    "Original Value": _display(d.original),
                    "Final Value": _display(d.final),
                    "Match Type": ev.type,
                    "Match Confidence": round(ev.score, 2),
                    "Cell Modified": "Yes" if modified else "No",
                    "Action Taken": action,
                    "Highlight Color": color,
                }
            )
    rows.sort(key=lambda r: (r["Check"], r["Respondent Row"], r["Source Cell"], r["Destination Cell"]))
    return rows


def build_review_rows(layout: SheetLayout, analysis: Analysis) -> List[dict]:
    agg: Dict[tuple, dict] = {}
    for rv in analysis.reviews:
        status = "Ambiguous" if rv.status == "ambiguous" else "Near match"
        for cand in rv.candidates:
            key = (rv.check_no, rv.src_title.strip(), cand.col)
            row = agg.get(key)
            if row is None:
                row = agg[key] = {
                    "Check": rv.check_no,
                    "Source Question": rv.src_q,
                    "Destination Question": rv.dst_q,
                    "Source Title": rv.src_title,
                    "Source Cells": [],
                    "Occurrences": 0,
                    "Destination Header": cand.header,
                    "Destination Header Cell": addr(layout.header_row, cand.col),
                    "Similarity Score": round(cand.result.score, 2),
                    "Candidate Type": cand.result.type,
                    "Status": status,
                    "Reason for Manual Review": (
                        rv.reason if rv.status == "ambiguous" else cand.result.reason
                    ),
                }
            row["Occurrences"] += 1
            a = addr(rv.row, rv.src_col)
            if len(row["Source Cells"]) < 10 and a not in row["Source Cells"]:
                row["Source Cells"].append(a)
    rows = []
    for key in sorted(agg, key=lambda k: (k[0], k[1].casefold(), k[2])):
        r = agg[key]
        cells = r["Source Cells"]
        more = "" if r["Occurrences"] <= len(cells) else ", ..."
        r["Source Cells"] = ", ".join(cells) + more
        rows.append(r)
    return rows


def build_statistics(analysis: Analysis, decisions: List[Decision], checks=None) -> Tuple[List[dict], dict]:
    checks = list(checks if checks is not None else CHECKS)
    per: Dict[int, Counter] = {n: Counter() for n, _, _ in checks}
    for n, c in analysis.titles_checked.items():
        per[n]["titles_checked"] = c
    for ev in analysis.events:
        per[ev.check_no][ev.type] += 1
    for rv in analysis.reviews:
        per[rv.check_no]["ambiguous" if rv.status == "ambiguous" else "near_match"] += 1
    for d in decisions:
        c = per[d.winner.check_no]
        c["unique_cells_matched"] += 1
        c["duplicates_suppressed"] += len(d.others)
        c[{ONE: "changed", TWO: "preserved_2", ZERO: "already_0", BLANK: "blank",
           FORMULA: "other", OTHER: "other"}[d.kind]] += 1
        if d.kind == FORMULA:
            c["formula_skipped"] += 1

    keys = [
        "titles_checked", EXACT, NORMALIZED, ALIAS, FUZZY, "ambiguous", "near_match",
        "unique_cells_matched", "changed", "preserved_2", "already_0", "blank", "other",
        "duplicates_suppressed", "formula_skipped",
    ]
    rows = []
    totals = Counter()
    for n, sq, dq in checks:
        row = {"check": n, "source": sq, "destination": dq}
        for k in keys:
            row[k] = int(per[n][k])
            totals[k] += int(per[n][k])
        row["review_total"] = row["ambiguous"] + row["near_match"]
        rows.append(row)
    tot = {k: int(totals[k]) for k in keys}
    tot["review_total"] = tot["ambiguous"] + tot["near_match"]
    tot["title_matches"] = tot[EXACT] + tot[NORMALIZED] + tot[ALIAS] + tot[FUZZY]
    return rows, tot


# --------------------------------------------------------------------------
# Step 3 + 4: write and verify
# --------------------------------------------------------------------------
def _fill(color_name: str) -> PatternFill:
    hexv = COLOR_HEX[color_name]
    return PatternFill(fill_type="solid", start_color=hexv, end_color=hexv)


def _rgb_tail(cell) -> str:
    try:
        rgb = cell.fill.fgColor.rgb
        return rgb[-6:].upper() if isinstance(rgb, str) else ""
    except Exception:  # pragma: no cover
        return ""


def _cell_signature(c) -> tuple:
    """(value, fill type, fill rgb, number format) - works for normal and read-only cells."""
    fill = c.fill
    ftype = getattr(fill, "fill_type", None) if fill is not None else None
    return (c.value, ftype, _rgb_tail(c), c.number_format or "General")


def _fingerprints(data: bytes, skip: Dict[Tuple[str, int, int], str]):
    """Stream one workbook (read-only mode, low memory).

    Returns sheet names, per-sheet (hash of every cell except ``skip``, dimensions) and the
    (value, fill) of the skipped cells.
    """
    wb = open_workbook(data, read_only=True)
    names = list(wb.sheetnames)
    sheets: Dict[str, Tuple[str, Tuple[int, int]]] = {}
    special: Dict[Tuple[str, int, int], tuple] = {}
    for ws in wb.worksheets:
        h = hashlib.sha256()
        n_rows = n_cols = 0
        for r, row in enumerate(ws.iter_rows(), start=1):
            n_rows = r
            for c_idx, c in enumerate(row, start=1):
                n_cols = max(n_cols, c_idx)
                value, ftype, rgb, numfmt = _cell_signature(c)
                key = (ws.title, r, c_idx)
                if key in skip:
                    special[key] = (value, rgb)
                    continue
                if value is None and ftype in (None, "none") and numfmt == "General":
                    continue  # empty and unstyled: identical whether stored or not
                h.update(repr((r, c_idx, value, ftype, rgb, numfmt)).encode())
        sheets[ws.title] = (h.hexdigest(), (n_rows, n_cols))
    wb.close()
    return names, sheets, special


def _detailed_diff(original: bytes, output: bytes, changed) -> List[str]:
    """Slow path used only after a fingerprint mismatch: list the differing cells."""
    issues: List[str] = []
    wo, wn = open_workbook(original), open_workbook(output)
    for name in wo.sheetnames:
        so, sn = wo[name], wn[name]
        for row in so.iter_rows():
            for co in row:
                if (name, co.row, co.column) in changed:
                    continue
                cn = sn.cell(row=co.row, column=co.column)
                if co.value != cn.value:
                    issues.append(f"{name}!{co.coordinate}: value changed unexpectedly")
                elif (co.fill.fill_type, _rgb_tail(co)) != (cn.fill.fill_type, _rgb_tail(cn)):
                    issues.append(f"{name}!{co.coordinate}: fill changed unexpectedly")
                elif co.number_format != cn.number_format:
                    issues.append(f"{name}!{co.coordinate}: number format changed unexpectedly")
                if len(issues) >= 50:
                    return issues + ["... more discrepancies not listed"]
    return issues


def verify_output(original: bytes, output: bytes, changed: Dict[Tuple[str, int, int], str]) -> List[str]:
    """Reopen both workbooks and prove that only the intended cells differ.

    Returns a list of discrepancies (empty list = verified).
    """
    issues: List[str] = []
    names_o, sheets_o, special_o = _fingerprints(original, changed)
    names_n, sheets_n, special_n = _fingerprints(output, changed)
    if names_o != names_n:
        return [f"Worksheet names/order differ: {names_o} vs {names_n}"]
    for key, color in changed.items():
        label = f"{key[0]}!{addr(key[1], key[2])}"
        if classify_value(special_o.get(key, (None,))[0]) != ONE:
            issues.append(f"{label}: changed although the original value was not 1")
        new_val, new_fill = special_n.get(key, (None, ""))
        if classify_value(new_val) != ZERO:
            issues.append(f"{label}: new value is not 0")
        if new_fill != COLOR_HEX[color]:
            issues.append(f"{label}: highlight colour missing")
    for name in names_o:
        if sheets_o[name][0] != sheets_n[name][0]:
            issues.extend(_detailed_diff(original, output, changed))
            break
    return issues[:51]


def process_workbook(
    data: bytes,
    filename: str,
    layout: SheetLayout,
    matcher: TitleMatcher,
    *,
    verify: bool = True,
    processed_at: Optional[datetime] = None,
    checks=None,
) -> ProcessResult:
    checks = list(checks if checks is not None else CHECKS)
    wb = open_workbook(data)  # the only full load: used for reading AND for writing the copy
    if layout.sheet not in wb.sheetnames:
        raise MappingError(f"Worksheet '{layout.sheet}' was not found in the workbook.")
    ws = wb[layout.sheet]
    snap = snapshot_from_sheet(ws, data)
    errors, layout_warnings = validate_layout(snap, layout, checks)
    if errors:
        raise MappingError(errors)

    analysis = analyze(snap, layout, matcher, checks)
    decisions = decide(snap, analysis)

    # --- safety net: never touch a source-question column --------------------------
    src_cols = {c for q in sources_of(checks) for c in layout.cols(q)}
    for d in decisions:
        if d.col in src_cols:  # pragma: no cover - validate_layout already forbids overlaps
            raise MappingError(f"Internal safety stop: {addr(d.row, d.col)} is a source-question cell.")

    changed: Dict[Tuple[str, int, int], str] = {}
    changed_cells: List[dict] = []
    for d in decisions:
        if d.kind != ONE:
            continue
        cell = ws.cell(row=d.row, column=d.col)
        cell.value = d.final
        cell.fill = _fill(d.color)
        changed[(layout.sheet, d.row, d.col)] = d.color
        changed_cells.append(
            {
                "cell": addr(d.row, d.col),
                "row": d.row,
                "col": d.col,
                "color": d.color,
                "match_type": d.winner.type,
                "source_title": d.winner.src_title,
                "header": d.winner.dst_header,
            }
        )
    buf = io.BytesIO()
    wb.save(buf)
    corrected = buf.getvalue()
    del wb, ws, snap  # free memory before verification re-opens the files
    gc.collect()

    issues: List[str] = []
    if verify:
        issues = verify_output(data, corrected, changed)

    audit_rows = build_audit_rows(layout, decisions)
    review_rows = build_review_rows(layout, analysis)
    per_check, totals = build_statistics(analysis, decisions, checks)

    return ProcessResult(
        filename=filename,
        sheet=layout.sheet,
        processed_at=processed_at or datetime.now(),
        threshold=matcher.threshold,
        layout=layout,
        corrected_bytes=corrected,
        audit_rows=audit_rows,
        review_rows=review_rows,
        per_check=per_check,
        totals=totals,
        changed_cells=changed_cells,
        feature_warnings=detect_unsupported_features(data),
        layout_warnings=layout_warnings,
        verification_issues=issues,
        verified=verify and not issues,
        scorer=scorer_name(),
        checks=checks,
    )


# --------------------------------------------------------------------------
# Preview (dry run - nothing is written)
# --------------------------------------------------------------------------
def preview_matches(data: bytes, layout: SheetLayout, matcher: TitleMatcher, checks=None) -> dict:
    """Dry run: distinct title matches, review items and how many cells would change."""
    checks = list(checks if checks is not None else CHECKS)
    snap = load_snapshot(data, layout.sheet)
    errors, warnings = validate_layout(snap, layout, checks)
    if errors:
        raise MappingError(errors)
    analysis = analyze(snap, layout, matcher, checks)
    decisions = decide(snap, analysis)
    _, totals = build_statistics(analysis, decisions, checks)

    agg: Dict[tuple, dict] = {}
    for ev in analysis.events:
        key = (ev.check_no, ev.src_title.strip().casefold(), ev.dst_col, ev.type)
        r = agg.get(key)
        if r is None:
            r = agg[key] = {
                "Check": f"{ev.check_no}: {ev.src_q} -> {ev.dst_q}",
                "Source Title": ev.src_title,
                "Destination Header": ev.dst_header,
                "Classification": ev.type,
                "Colour if changed": COLOR_FOR[ev.type],
                "Score": round(ev.score, 1),
                "Occurrences": 0,
            }
        r["Occurrences"] += 1
    matches = sorted(agg.values(), key=lambda r: (r["Check"], -PRIORITY[r["Classification"]], r["Source Title"].casefold()))
    review = build_review_rows(layout, analysis)
    return {
        "matches": matches,
        "review": review,
        "totals": totals,
        "would_change": totals["changed"],
        "warnings": warnings,
    }

# ======================================================================
# module: report
# ======================================================================
"""Excel audit report (Summary / Audit Log / Review Matches)."""


import io
from typing import List

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter










HEADER_FILL = PatternFill("solid", start_color="1F3864", end_color="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
SECTION_FONT = Font(bold=True, size=12, color="1F3864")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

AUDIT_COLUMNS = [
    "Worksheet", "Check", "Source Question", "Destination Question", "Respondent Row",
    "Source Title", "Source Cell", "Destination Title Header", "Destination Header Cell",
    "Destination Cell", "Original Value", "Final Value", "Match Type", "Match Confidence",
    "Cell Modified", "Action Taken", "Highlight Color",
]
REVIEW_COLUMNS = [
    "Check", "Source Question", "Destination Question", "Source Title", "Source Cells",
    "Occurrences", "Destination Header", "Destination Header Cell", "Similarity Score",
    "Candidate Type", "Status", "Reason for Manual Review",
]

PER_CHECK_COLUMNS = [
    ("check", "Check"), ("source", "Source Q"), ("destination", "Destination Q"),
    ("titles_checked", "Source titles checked"),
    ("exact", "Exact matches"), ("normalized", "Normalized matches"),
    ("alias", "Recognised-variant (alias) matches"), ("fuzzy", "Fuzzy matches"),
    ("changed", "Cells changed 1 to 0"), ("preserved_2", "Cells preserved (value 2)"),
    ("already_0", "Cells already 0"), ("blank", "Blank cells skipped"),
    ("other", "Other values skipped"), ("duplicates_suppressed", "Duplicate matches suppressed"),
    ("ambiguous", "Ambiguous (manual review)"), ("near_match", "Near matches (manual review)"),
]


def _style_header(ws, row: int, ncols: int):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BORDER


def _write_summary(ws, res: ProcessResult) -> None:
    ws.sheet_view.showGridLines = False
    r = 1
    ws.cell(row=r, column=1, value="OTT Title Matching - Audit Report").font = Font(bold=True, size=16, color="1F3864")
    r += 2

    info = [
        ("Original filename", res.filename),
        ("Processing date and time", res.processed_at.strftime("%Y-%m-%d %H:%M:%S")),
        ("Selected worksheet", res.sheet),
        ("Fuzzy-matching threshold", f"{res.threshold:g}%"),
        ("Similarity scorer", res.scorer or scorer_name()),
        ("Question-label row", res.layout.label_row),
        ("Title-header row", res.layout.header_row),
        ("First respondent row", res.layout.first_row),
        ("Last respondent row", res.layout.last_row),
        ("Output verification", "PASSED - re-opened the saved workbook; only intended cells differ"
         if res.verified else "NOT VERIFIED / issues found: " + "; ".join(res.verification_issues[:5])),
    ]
    for k, v in info:
        ws.cell(row=r, column=1, value=k).font = Font(bold=True)
        ws.cell(row=r, column=2, value=v)
        r += 1
    r += 1

    ws.cell(row=r, column=1, value="Colour legend").font = SECTION_FONT
    r += 1
    legend = [
        ("Yellow", "Exact match, changed from 1 to 0."),
        ("Green", "Normalized or fuzzy title match (incl. recognised variants), changed from 1 to 0."),
    ]
    for name, text in legend:
        c = ws.cell(row=r, column=1, value=name)
        c.fill = PatternFill("solid", start_color=COLOR_HEX[name], end_color=COLOR_HEX[name])
        c.font = Font(bold=True, color="000000" if name == "Yellow" else "FFFFFF")
        c.border = BORDER
        ws.cell(row=r, column=2, value=f"{name} = {text}")
        r += 1
    r += 1

    ws.cell(row=r, column=1, value="Question-pair checks used").font = SECTION_FONT
    r += 1
    for n, s_, d_ in res.checks:
        ws.cell(row=r, column=1, value=f"Check {n}")
        ws.cell(row=r, column=2, value=f"{s_} (source) -> {d_} (destination)")
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Question mappings used").font = SECTION_FONT
    r += 1
    for j, h in enumerate(["Question", "Role", "Columns", "Label cell"], start=1):
        ws.cell(row=r, column=j, value=h)
    _style_header(ws, r, 4)
    r += 1
    sources = set(sources_of(res.checks))
    for qid in questions_of(res.checks):
        q = res.layout.questions.get(qid)
        role = "Source (never modified)" if qid in sources else "Destination (1 -> 0 corrections)"
        vals = [qid, role, columns_to_spec(q.columns) if q else "", (q.label_cell if q else "") or "(manual)"]
        for j, v in enumerate(vals, start=1):
            ws.cell(row=r, column=j, value=v).border = BORDER
        r += 1
    r += 1

    ws.cell(row=r, column=1, value="Overall totals").font = SECTION_FONT
    r += 1
    t = res.totals
    overall = [
        ("Source titles checked", t["titles_checked"]),
        ("Exact matches", t["exact"]),
        ("Normalized matches", t["normalized"]),
        ("Recognised-variant (alias) matches", t["alias"]),
        ("Fuzzy matches", t["fuzzy"]),
        ("Cells changed from 1 to 0", t["changed"]),
        ("Cells preserved because value was 2", t["preserved_2"]),
        ("Cells already containing 0", t["already_0"]),
        ("Blank cells skipped", t["blank"]),
        ("Other values skipped (incl. formulas)", t["other"]),
        ("Duplicate matches suppressed (cell already handled)", t["duplicates_suppressed"]),
        ("Ambiguous matches requiring review", t["ambiguous"]),
        ("Near matches requiring review", t["near_match"]),
    ]
    for k, v in overall:
        ws.cell(row=r, column=1, value=k).border = BORDER
        c = ws.cell(row=r, column=2, value=v)
        c.border = BORDER
        c.alignment = Alignment(horizontal="left")
        r += 1
    r += 1

    ws.cell(row=r, column=1, value="Statistics per question-pair check").font = SECTION_FONT
    r += 1
    for j, (_, title) in enumerate(PER_CHECK_COLUMNS, start=1):
        ws.cell(row=r, column=j, value=title)
    _style_header(ws, r, len(PER_CHECK_COLUMNS))
    ws.row_dimensions[r].height = 45
    r += 1
    for row in res.per_check:
        for j, (key, _) in enumerate(PER_CHECK_COLUMNS, start=1):
            ws.cell(row=r, column=j, value=row[key]).border = BORDER
        r += 1
    for j, (key, _) in enumerate(PER_CHECK_COLUMNS, start=1):
        if j == 1:
            v = "Total"
        elif key in ("source", "destination"):
            v = ""
        else:
            v = res.totals[key]
        c = ws.cell(row=r, column=j, value=v)
        c.font = Font(bold=True)
        c.border = BORDER
    r += 2

    ws.cell(row=r, column=1, value="Rules applied").font = SECTION_FONT
    r += 1
    notes = [
        "Destination cell = 1 -> changed to 0 and highlighted. = 2 -> preserved. = 0, blank or any other value -> left unchanged.",
        "Source-question answers and unrelated cells are never modified.",
        "Ambiguous or low-confidence matches are listed in 'Review Matches' and are NOT applied.",
        PRIORITY_RULE_TEXT,
        "Match counts above count title matches (before de-duplication); cell counts are per unique destination cell.",
    ]
    for n in notes:
        c = ws.cell(row=r, column=1, value=n)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
        ws.row_dimensions[r].height = 15 * (1 + len(n) // 120)
        r += 1

    if res.feature_warnings or res.layout_warnings:
        r += 1
        ws.cell(row=r, column=1, value="Warnings").font = SECTION_FONT
        r += 1
        for w in [*res.feature_warnings, *res.layout_warnings]:
            ws.cell(row=r, column=1, value=w)
            r += 1

    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 34
    for j in range(3, len(PER_CHECK_COLUMNS) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 16


def _format_table(ws, ncols: int, widths: List[int]) -> None:
    _style_header(ws, 1, ncols)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[1].height = 32


def build_audit_report(res: ProcessResult) -> bytes:
    """Return the audit report workbook as bytes."""
    audit_df = pd.DataFrame(res.audit_rows, columns=AUDIT_COLUMNS)
    review_df = pd.DataFrame(res.review_rows, columns=REVIEW_COLUMNS)

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        audit_df.to_excel(writer, sheet_name="Audit Log", index=False)
        review_df.to_excel(writer, sheet_name="Review Matches", index=False)
        wb = writer.book
        summary = wb.create_sheet("Summary", 0)
        _write_summary(summary, res)

        ws = wb["Audit Log"]
        _format_table(ws, len(AUDIT_COLUMNS), [14, 7, 10, 12, 11, 24, 11, 24, 12, 12, 10, 10, 12, 11, 10, 46, 12])
        color_col = AUDIT_COLUMNS.index("Highlight Color") + 1
        for row in range(2, ws.max_row + 1):
            cell = ws.cell(row=row, column=color_col)
            if cell.value in COLOR_HEX:
                cell.fill = PatternFill("solid", start_color=COLOR_HEX[cell.value], end_color=COLOR_HEX[cell.value])

        ws = wb["Review Matches"]
        _format_table(ws, len(REVIEW_COLUMNS), [7, 10, 12, 26, 30, 11, 26, 14, 11, 12, 12, 60])
        if review_df.empty:
            ws.cell(row=2, column=1, value="No ambiguous or near-matching titles needed review.")
        wb.active = 0
    return buf.getvalue()

